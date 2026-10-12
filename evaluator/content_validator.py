from __future__ import annotations

import csv
import json
import re
from dataclasses import dataclass, field
from io import StringIO
from pathlib import Path
from typing import Any

from .promise_parser import FilePromise, SectionConstraint


@dataclass
class FileValidationResult:
    filename: str
    file_type: str
    exists: bool = False
    valid: bool = False
    score: int = 0
    max_score: int = 100
    details: list[str] = field(default_factory=list)
    section_results: list[dict[str, Any]] = field(default_factory=list)


class ContentValidator:
    def validate_file(self, file_path: Path, promise: FilePromise) -> FileValidationResult:
        result = FileValidationResult(filename=promise.filename, file_type=promise.file_type)

        if not file_path.exists():
            result.exists = False
            result.valid = False
            result.score = 0
            result.details.append(f"文件不存在: {promise.filename}")
            return result

        result.exists = True
        content = self._read_file(file_path)
        if content is None:
            result.details.append(f"文件无法读取: {promise.filename}")
            result.score = 20
            return result

        if len(content.strip()) < 10:
            result.details.append("文件内容过少 (<10字符)")
            result.score = 30
            return result

        validator = self._get_validator(promise.file_type)
        result = validator(content, file_path, promise)

        return result

    def _read_file(self, path: Path) -> str | None:
        try:
            return path.read_text(encoding="utf-8", errors="replace")
        except Exception:
            try:
                return path.read_text(encoding="gbk", errors="replace")
            except Exception:
                return None

    def _get_validator(self, file_type: str):
        validators = {
            "md": self._validate_md,
            "markdown": self._validate_md,
            "csv": self._validate_csv,
            "json": self._validate_json,
            "xlsx": self._validate_xlsx,
            "xls": self._validate_xlsx,
            "sql": self._validate_sql,
            "txt": self._validate_txt,
            "py": self._validate_code,
            "js": self._validate_code,
            "ts": self._validate_code,
        }
        return validators.get(file_type, self._validate_generic)

    def _validate_md(self, content: str, path: Path, promise: FilePromise) -> FileValidationResult:
        result = FileValidationResult(filename=promise.filename, file_type=promise.file_type, exists=True)

        headers = self._parse_md_headers(content)
        sections = self._parse_md_sections(content, headers)

        base_score = 20
        result.details.append(f"文件大小: {len(content)}字符, {content.count(chr(10))}行")
        result.details.append(f"标题数量: {len(headers)}")

        for section in promise.sections:
            section_result = {"name": section.name, "passed": False, "detail": ""}

            matching_header = None
            for h in headers:
                if section.name.lower() in h["text"].lower() or h["text"].lower() in section.name.lower():
                    matching_header = h
                    break

            if not matching_header and not section.must_exist:
                section_result["passed"] = True
                section_result["detail"] = "可选章节，未找到"
                result.section_results.append(section_result)
                continue

            if not matching_header:
                section_result["detail"] = f"缺少章节: {section.name}"
                section_result["passed"] = False
            else:
                section_text = sections.get(matching_header["text"], "")
                section_result["detail"] = f"找到章节: {section.name}"

                if section.min_items is not None:
                    items = re.findall(r'^\s*[-*+]\s+', section_text, re.MULTILINE)
                    numbered_items = re.findall(r'^\s*\d+[\.、)\s]', section_text, re.MULTILINE)
                    count = len(items) + len(numbered_items)
                    if count < section.min_items:
                        section_result["detail"] += f", {section.name}仅{count}条(要求≥{section.min_items})"
                        section_result["passed"] = False
                    else:
                        section_result["detail"] += f", {section.name}{count}条(≥{section.min_items} ✓)"
                        section_result["passed"] = True

                if section.min_length is not None:
                    length = len(section_text.strip())
                    if length < section.min_length:
                        section_result["detail"] += f", 长度{length}(要求≥{section.min_length})"
                        section_result["passed"] = False

                if section.max_length is not None:
                    length = len(section_text.strip())
                    if length > section.max_length:
                        section_result["detail"] += f", 长度{length}(要求≤{section.max_length})"
                        section_result["passed"] = False

                if not section_result.get("passed") and section_result["detail"].startswith("找到章节"):
                    section_result["passed"] = True

            base_score += 15 if section_result["passed"] else 0
            result.section_results.append(section_result)

        if not promise.sections:
            if len(content.strip()) > 50:
                base_score = 70
                result.details.append("无明确章节要求，内容充实")
            else:
                base_score = 50
                result.details.append("无明确章节要求，内容较少")

        result.score = min(100, base_score)
        result.valid = result.score >= 60
        return result

    def _validate_csv(self, content: str, path: Path, promise: FilePromise) -> FileValidationResult:
        result = FileValidationResult(filename=promise.filename, file_type=promise.file_type, exists=True)

        try:
            reader = csv.DictReader(StringIO(content))
            headers = reader.fieldnames or []
            rows = list(reader)
            row_count = len(rows)
        except Exception as e:
            result.details.append(f"CSV解析失败: {e}")
            result.score = 30
            return result

        result.details.append(f"列数: {len(headers)}, 行数: {row_count}")
        result.details.append(f"列名: {', '.join(headers[:5])}")

        base_score = 30
        if row_count > 0:
            base_score += 20
        if len(headers) >= 2:
            base_score += 10

        for section in promise.sections:
            if section.name.lower() in [h.lower() for h in headers]:
                base_score += 10
                result.details.append(f"找到列: {section.name}")
            if section.min_items is not None and row_count >= section.min_items:
                base_score += 10
                result.details.append(f"行数{row_count}≥{section.min_items}")

        result.score = min(100, base_score)
        result.valid = result.score >= 50
        return result

    def _validate_json(self, content: str, path: Path, promise: FilePromise) -> FileValidationResult:
        result = FileValidationResult(filename=promise.filename, file_type=promise.file_type, exists=True)

        try:
            data = json.loads(content)
        except json.JSONDecodeError as e:
            result.details.append(f"JSON解析失败: {e}")
            result.score = 20
            return result

        base_score = 30

        if isinstance(data, dict):
            keys = list(data.keys())
            result.details.append(f"类型: object, 键数: {len(keys)}")
            result.details.append(f"键名: {', '.join(keys[:8])}")
            base_score += 20 if len(keys) > 0 else 0

            for section in promise.sections:
                matching_key = None
                for k in keys:
                    if section.name.lower() in k.lower() or k.lower() in section.name.lower():
                        matching_key = k
                        break
                if matching_key:
                    base_score += 10
                    result.details.append(f"找到键: {section.name}")
                    val = data[matching_key]
                    if isinstance(val, list) and section.min_items is not None:
                        if len(val) >= section.min_items:
                            result.details.append(f"数组长度{len(val)}≥{section.min_items}")
                        else:
                            result.details.append(f"数组长度{len(val)}<{section.min_items}")

        elif isinstance(data, list):
            result.details.append(f"类型: array, 长度: {len(data)}")
            base_score += 20
            for section in promise.sections:
                if section.min_items is not None and len(data) >= section.min_items:
                    base_score += 15
                    result.details.append(f"数组长度{len(data)}≥{section.min_items}")
        else:
            result.details.append(f"类型: {type(data).__name__}")
            base_score += 10

        result.score = min(100, base_score)
        result.valid = result.score >= 50
        return result

    def _validate_xlsx(self, content: str, path: Path, promise: FilePromise) -> FileValidationResult:
        result = FileValidationResult(filename=promise.filename, file_type=promise.file_type, exists=True)

        try:
            import openpyxl
            wb = openpyxl.load_workbook(path)
            sheet_names = wb.sheetnames
            sheet_count = len(sheet_names)
            result.details.append(f"Sheet数: {sheet_count}, 名称: {', '.join(sheet_names)}")

            base_score = 30
            if sheet_count > 0:
                base_score += 20

            for section in promise.sections:
                matching_sheet = None
                for s in sheet_names:
                    if section.name.lower() in s.lower() or s.lower() in section.name.lower():
                        matching_sheet = s
                        break
                if matching_sheet:
                    ws = wb[matching_sheet]
                    row_count = ws.max_row
                    col_count = ws.max_column
                    base_score += 15
                    result.details.append(f"Sheet '{section.name}': {row_count}行, {col_count}列")
                    if section.min_items is not None and row_count >= section.min_items:
                        result.details.append(f"行数{row_count}≥{section.min_items}")

            result.score = min(100, base_score)
            wb.close()
        except ImportError:
            result.details.append("openpyxl未安装，使用基础检查")
            file_size = path.stat().st_size
            if file_size > 100:
                result.score = 60
                result.details.append(f"文件大小: {file_size}字节")
            else:
                result.score = 40
                result.details.append(f"文件过小: {file_size}字节")
        except Exception as e:
            result.details.append(f"Excel解析失败: {e}")
            result.score = 30

        result.valid = result.score >= 50
        return result

    def _validate_sql(self, content: str, path: Path, promise: FilePromise) -> FileValidationResult:
        result = FileValidationResult(filename=promise.filename, file_type=promise.file_type, exists=True)

        statements = [s.strip() for s in content.split(';') if s.strip()]
        stmt_count = len(statements)

        result.details.append(f"语句数: {stmt_count}")

        base_score = 30
        if stmt_count > 0:
            base_score += 20

        has_create = bool(re.search(r'CREATE\s+(TABLE|DATABASE|INDEX)', content, re.IGNORECASE))
        has_insert = bool(re.search(r'INSERT\s+INTO', content, re.IGNORECASE))
        has_select = bool(re.search(r'SELECT\s+', content, re.IGNORECASE))

        if has_create:
            base_score += 10
            result.details.append("包含CREATE语句")
        if has_insert:
            base_score += 10
            result.details.append("包含INSERT语句")
        if has_select:
            base_score += 5
            result.details.append("包含SELECT语句")

        result.score = min(100, base_score)
        result.valid = result.score >= 50
        return result

    def _validate_code(self, content: str, path: Path, promise: FilePromise) -> FileValidationResult:
        result = FileValidationResult(filename=promise.filename, file_type=promise.file_type, exists=True)

        has_func = bool(re.search(r'\bdef\s+\w+|function\s+\w+|class\s+\w+', content))
        has_import = bool(re.search(r'\bimport\s+|from\s+\w+\s+import', content))
        line_count = content.count('\n') + 1

        result.details.append(f"行数: {line_count}")
        if has_func:
            result.details.append("包含函数/类定义")
        if has_import:
            result.details.append("包含import语句")

        base_score = 40
        if line_count > 10:
            base_score += 20
        if has_func:
            base_score += 20
        if has_import:
            base_score += 10

        result.score = min(100, base_score)
        result.valid = result.score >= 50
        return result

    def _validate_txt(self, content: str, path: Path, promise: FilePromise) -> FileValidationResult:
        result = FileValidationResult(filename=promise.filename, file_type=promise.file_type, exists=True)

        line_count = content.count('\n') + 1
        word_count = len(content.split())

        result.details.append(f"行数: {line_count}, 词数: {word_count}")

        base_score = 30
        if line_count > 5:
            base_score += 20
        if word_count > 50:
            base_score += 20

        result.score = min(100, base_score)
        result.valid = result.score >= 50
        return result

    def _validate_generic(self, content: str, path: Path, promise: FilePromise) -> FileValidationResult:
        result = FileValidationResult(filename=promise.filename, file_type=promise.file_type, exists=True)

        file_size = path.stat().st_size
        result.details.append(f"文件大小: {file_size}字节")

        base_score = 40
        if file_size > 100:
            base_score += 30
        if len(content.strip()) > 50:
            base_score += 20

        result.score = min(100, base_score)
        result.valid = result.score >= 50
        return result

    def _parse_md_headers(self, content: str) -> list[dict[str, Any]]:
        headers = []
        for match in re.finditer(r'^(#{1,6})\s+(.+)$', content, re.MULTILINE):
            headers.append({
                "level": len(match.group(1)),
                "text": match.group(2).strip(),
                "pos": match.start(),
            })
        return headers

    def _parse_md_sections(self, content: str, headers: list[dict]) -> dict[str, str]:
        sections = {}
        for i, h in enumerate(headers):
            start = h["pos"]
            end = headers[i + 1]["pos"] if i + 1 < len(headers) else len(content)
            section_text = content[start:end]
            sections[h["text"]] = section_text
            for alias in _get_section_aliases(h["text"]):
                sections[alias] = section_text
        return sections


_SECTION_ALIASES = {
    "执行摘要": ["摘要", "执行摘要", "summary"],
    "关键发现": ["发现", "关键发现", "findings"],
    "风险提示": ["风险", "风险提示", "risks"],
    "行动建议": ["建议", "行动建议", "recommendations", "actions"],
    "数据质量": ["数据质量", "data quality"],
    "数据校验": ["数据校验", "data validation"],
    "趋势分析": ["趋势分析", "trend analysis"],
}

def _get_section_aliases(text: str) -> list[str]:
    for key, aliases in _SECTION_ALIASES.items():
        if text in aliases:
            return [key]
    return [text]
